import sqlite3
import subprocess
import sys

import pytest
from docx import Document
from sqlalchemy import create_engine

from db_export.cli import connection_url, main, parser
from db_export.document import build_document
from db_export.schema import read_schema


@pytest.fixture
def database(tmp_path):
    path = tmp_path / "sample.db"
    with sqlite3.connect(path) as db:
        db.executescript('''
            CREATE TABLE parent (tenant INTEGER, id INTEGER, PRIMARY KEY (tenant, id));
            CREATE TABLE child (
                id INTEGER PRIMARY KEY, tenant INTEGER NOT NULL, parent_id INTEGER,
                "中文名稱" VARCHAR(50) NOT NULL DEFAULT '',
                score NUMERIC(10, 2) DEFAULT 0 CHECK (score >= 0),
                UNIQUE (tenant, "中文名稱"),
                FOREIGN KEY (tenant, parent_id) REFERENCES parent (tenant, id)
            );
            CREATE INDEX ix_child_parent ON child(tenant, parent_id);
            CREATE TABLE ignored (id INTEGER);
        ''')
    return path


def test_full_cli_export(database, tmp_path):
    output = tmp_path / "nested" / "schema.docx"
    result = subprocess.run([sys.executable, "-m", "db_export", "--db-type", "sqlite",
                             "--database", str(database), "--exclude", "ignored", "-o", str(output)],
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    doc = Document(output)
    content = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "主鍵順序" not in content and "主键顺序" not in content
    assert "数据库结构说明书" in content
    assert "字段名称" in cells and "主键" in cells
    assert "中文名稱" in cells
    assert "VARCHAR(50)" in cells
    assert "NUMERIC(10, 2)" in cells
    assert "tenant, parent_id → parent (tenant, id)" in cells
    assert "ix_child_parent" in cells
    assert "score >= 0" in cells
    assert "ignored" not in cells
    assert all(t.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for t in doc.tables)


def test_existing_output_and_force(database, tmp_path):
    output = tmp_path / "schema.docx"
    output.write_bytes(b"existing")
    args = ["--db-type", "sqlite", "--database", str(database), "-o", str(output)]
    assert main(args) == 2
    assert output.read_bytes() == b"existing"
    assert main(args + ["--force"]) == 0
    Document(output)


def test_filter_failure_does_not_create_output(database, tmp_path):
    output = tmp_path / "schema.docx"
    assert main(["--db-type", "sqlite", "--database", str(database), "--include", "typo*",
                 "-o", str(output)]) == 2
    assert not output.exists()


def test_missing_sqlite_is_not_created(tmp_path):
    path = tmp_path / "missing.db"
    assert main(["--db-type", "sqlite", "--database", str(path)]) == 2
    assert not path.exists()


def test_password_is_not_url_encoded_twice():
    args = parser().parse_args(["--db-type", "mysql", "--user", "u", "--password", "p@ss:/?#",
                                "--database", "app"])
    url = connection_url(args)
    assert url.password == "p@ss:/?#"
    assert url.port == 3306


def test_malformed_url_does_not_echo_secret(capsys):
    assert main(["--url", "top-secret-not-a-url"]) == 2
    assert "top-secret" not in capsys.readouterr().err


def test_url_and_components_are_rejected():
    args = parser().parse_args(["--url", "sqlite:///a.db", "--host", "localhost"])
    with pytest.raises(ValueError, match="不可"):
        connection_url(args)


def test_comments_and_generated_columns(database):
    engine = create_engine(f"sqlite:///{database}")
    try:
        with engine.connect() as connection:
            tables, notices = read_schema(connection, includes=["child"])
        tables[0]["comment"] = "訂單明細"
        tables[0]["columns"][0].update(comment="識別碼", identity={"start": 1})
        doc = build_document(tables, "測試", "sample", None, notices)
        assert any("訂單明細" in p.text for p in doc.paragraphs)
        assert any("識別碼\nIdentity" in c.text for t in doc.tables for r in t.rows for c in r.cells)
    finally:
        engine.dispose()


@pytest.mark.parametrize("stage", ["connect", "read"])
def test_database_error_shows_safe_reason(stage, monkeypatch, capsys, tmp_path):
    from unittest.mock import MagicMock
    from sqlalchemy.exc import OperationalError
    from urllib.parse import quote
    import db_export.cli as cli

    secret = "test@secret:/123"
    original = Exception(1045, f"Access denied password={secret} encoded={quote(secret, safe='')}\nSQLSTATE 28000")
    error = OperationalError("PRIVATE SQL TEXT", {"secret": secret}, original)
    engine = MagicMock()
    monkeypatch.setattr(cli, "create_engine", lambda *a, **kw: engine)
    if stage == "connect":
        engine.connect.side_effect = error
    else:
        monkeypatch.setattr(cli, "read_schema", lambda *a, **kw: (_ for _ in ()).throw(error))
    result = main(["--db-type", "mysql", "--user", "reader", "--password", secret,
                   "--database", "app", "-o", str(tmp_path / "out.docx")])
    output = capsys.readouterr().err
    assert result == 1
    assert "1045" in output and "Access denied" in output and "28000" in output
    assert ("連接資料庫" if stage == "connect" else "讀取表結構") in output
    assert secret not in output and quote(secret, safe="") not in output
    assert "PRIVATE SQL TEXT" not in output
    assert "<REDACTED>" in output
    assert not (tmp_path / "out.docx").exists()
    engine.dispose.assert_called_once()

import pytest
from docx.oxml.ns import qn

from db_export.cli import parser
from db_export.document import build_document


def sample():
    return [{"name": "orders", "comment": "订单", "columns": [
        {"name": "id", "type": "BIGINT", "nullable": False},
        {"name": "note", "type": "VARCHAR(200)", "nullable": True}],
        "pk": ["id"], "indexes": [], "foreign_keys": [], "unique": [], "checks": []}]


@pytest.mark.parametrize("theme,fill,page_break", [
    ("classic", "243F60", True), ("compact", "243F60", False), ("monochrome", "E6E6E6", True)])
def test_theme_navigation_pagination_and_footer(theme, fill, page_break):
    doc = build_document(sample(), "数据库结构说明书", "app", "public", theme=theme)
    header = doc.tables[1].rows[0]
    assert header.cells[0]._tc.xpath("./w:tcPr/w:shd")[0].get(qn("w:fill")) == fill
    assert header._tr.xpath("./w:trPr/w:tblHeader")
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in doc.tables[1].rows)
    anchors = doc.element.xpath("//w:hyperlink/@w:anchor")
    targets = doc.element.xpath("//w:bookmarkStart/@w:name")
    assert anchors == targets == ["table_1"]
    heading = next(p for p in doc.paragraphs if p.text == "1  orders")
    assert heading.paragraph_format.page_break_before is page_break
    assert doc.tables[0].cell(1, 2).text == "2"
    footer = doc.sections[0].footer._element
    assert footer.xpath(".//w:fldSimple/@w:instr") == ["PAGE", "NUMPAGES"]
    assert "主键顺序" not in doc.element.xml
    assert doc.tables[1].cell(1, 2).paragraphs[0].runs[0].font.name == "Consolas"


@pytest.mark.parametrize("theme,override", [("compact", True), ("classic", False)])
def test_explicit_page_break_overrides_theme(theme, override):
    doc = build_document(sample(), "标题", "app", None, theme=theme, table_page_break=override)
    heading = next(p for p in doc.paragraphs if p.text == "1  orders")
    assert heading.paragraph_format.page_break_before is override


def test_cli_accepts_styles_and_page_break_flags():
    args = parser().parse_args(["--db-type", "sqlite", "--theme", "compact", "--table-page-break"])
    assert args.theme == "compact" and args.table_page_break is True
    args = parser().parse_args(["--db-type", "sqlite", "--no-table-page-break"])
    assert args.table_page_break is False

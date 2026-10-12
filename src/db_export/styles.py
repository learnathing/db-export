"""Reusable typography and OOXML helpers for the Word exporter."""

from dataclasses import dataclass

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


@dataclass(frozen=True)
class Theme:
    accent: str
    header: str
    header_text: str
    stripe: str
    border: str
    font_size: float = 9
    padding: int = 80
    spacing: int = 3
    page_break: bool = True


THEMES = {
    "classic": Theme("243F60", "243F60", "FFFFFF", "F3F6FA", "D6DFE8"),
    "compact": Theme("243F60", "243F60", "FFFFFF", "F3F6FA", "D6DFE8",
                     font_size=8.5, padding=45, spacing=1, page_break=False),
    "monochrome": Theme("202020", "E6E6E6", "000000", "F6F6F6", "B5B5B5"),
}


def element(tag, **attributes):
    node = OxmlElement("w:" + tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    return node


def set_font(run, *, color=None, size=None, mono=False):
    run.font.name = "Consolas" if mono else "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def configure_document(doc, theme, title):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.5)
    section.footer_distance = Cm(0.7)
    for name, size in (("Normal", 10), ("Title", 26), ("Heading 1", 16), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(theme.accent if name.startswith("Heading") else "202020")
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        fmt = style.paragraph_format
        fmt.space_before = Pt(12 if name.startswith("Heading") else 0)
        fmt.space_after = Pt(6)
        fmt.line_spacing = 1.1
        fmt.widow_control = True
        fmt.keep_with_next = name != "Normal"
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.title = title
    doc.core_properties.subject = "数据库结构与字段说明"
    doc.settings.element.append(element("updateFields", val="true"))
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(26.7), WD_TAB_ALIGNMENT.RIGHT)
    short_title = " ".join(title.split())
    short_title = short_title[:32] + ("…" if len(short_title) > 32 else "")
    footer.add_run(short_title + "\t第 ")
    for field, suffix in (("PAGE", " 页 / 共 "), ("NUMPAGES", " 页")):
        footer._p.append(element("fldSimple", instr=field))
        footer.add_run(suffix)
    for run in footer.runs:
        set_font(run, size=8, color="666666")


def style_table(table, theme, mono_columns=(), center_columns=()):
    props = table._tbl.tblPr
    borders = element("tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element(edge, val="single", sz=4, color=theme.border))
    props.append(borders)
    margins = element("tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        margins.append(element(side, w=theme.padding, type="dxa"))
    props.append(margins)
    for index, row in enumerate(table.rows):
        # Exceptionally long notes must be allowed to flow over a page.
        if max(len(c.text) for c in row.cells) < 1500:
            row._tr.get_or_add_trPr().append(element("cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(element("tblHeader"))
        fill = theme.header if index == 0 else theme.stripe if index % 2 == 0 else "FFFFFF"
        for col, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(element("shd", fill=fill))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(theme.spacing)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.keep_with_next = index == 0
                p.paragraph_format.line_spacing = 1.05
                if col in center_columns:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.bold = index == 0
                    set_font(run, size=theme.font_size, mono=index > 0 and col in mono_columns,
                             color=theme.header_text if index == 0 else "202020")


def bookmark(paragraph, number):
    paragraph._p.append(element("bookmarkStart", id=number, name=f"table_{number}"))
    paragraph._p.append(element("bookmarkEnd", id=number))


def link_cell(cell, number, theme):
    paragraph = cell.paragraphs[0]
    label = paragraph.text
    paragraph.clear()
    link = element("hyperlink", anchor=f"table_{number}", history=1)
    run = element("r")
    props = element("rPr")
    props.append(element("color", val=theme.accent))
    props.append(element("u", val="single"))
    props.append(element("sz", val=int(theme.font_size * 2)))
    props.append(element("rFonts", ascii="Consolas", hAnsi="Consolas", eastAsia="Microsoft YaHei"))
    run.append(props)
    content = element("t")
    content.text = label
    run.append(content)
    link.append(run)
    paragraph._p.append(link)

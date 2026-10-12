"""Word layout independent from database access."""

import os
import re
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from .styles import THEMES, bookmark, configure_document, link_cell, style_table


def text(value):
    if value is None:
        return "—"
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def grid(doc, headers, rows, widths, theme, *, mono_columns=(), center_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for row, values in [(table.rows[0], headers)] + [(table.add_row(), values) for values in rows]:
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.text = text(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = False
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    style_table(table, theme, mono_columns, center_columns)
    return table


def build_document(tables, title, database, schema, warnings=(), *, theme="classic", table_page_break=None):
    palette = THEMES[theme]
    page_break = palette.page_break if table_page_break is None else table_page_break
    doc = Document()
    configure_document(doc, palette, text(title))
    doc.add_heading(text(title), 0)
    doc.add_paragraph(f"数据库：{text(database)}    Schema：{text(schema or '默认')}\n"
                      f"导出时间：{datetime.now().astimezone():%Y-%m-%d %H:%M:%S %z}    数据表数量：{len(tables)}")
    doc.add_paragraph("本文档记录数据表结构、字段说明及约束，供开发与维护查阅。类型中的括号保留数据库返回的长度或精度；未提供的说明以 — 表示。")
    for warning in warnings:
        doc.add_paragraph(warning)
    doc.add_heading("数据表总览", 1)
    overview = grid(doc, ["序号", "数据表", "字段数", "说明"],
                    [(i, t["name"], len(t["columns"]), t["comment"]) for i, t in enumerate(tables, 1)],
                    [1.2, 7, 1.5, 17], palette, center_columns=(0, 2))
    for i, row in enumerate(overview.rows[1:], 1):
        link_cell(row.cells[1], i, palette)
    for i, table in enumerate(tables, 1):
        heading = doc.add_heading(f"{i}  {text(table['name'])}", 1)
        heading.paragraph_format.page_break_before = page_break
        bookmark(heading, i)
        description = doc.add_paragraph(f"表说明：{text(table['comment'])}")
        description.paragraph_format.keep_with_next = True
        rows = []
        for j, c in enumerate(table["columns"], 1):
            extras = []
            if c.get("identity"):
                extras.append(f"Identity：{c['identity']}")
            elif c.get("autoincrement") is True:
                extras.append("自增")
            if c.get("computed"):
                extras.append(f"计算表达式：{c['computed'].get('sqltext', '')}")
            description = "\n".join(([c["comment"]] if c.get("comment") else []) + extras) or None
            rows.append((j, c["name"], c["type"], "是" if c["name"] in table["pk"] else "否",
                         "是" if c.get("nullable") else "否", c.get("default"), description))
        grid(doc, ["序号", "字段名称", "类型／长度", "主键", "允许空", "默认值", "说明"],
             rows, [1.1, 5, 3.5, 1.1, 1.5, 4, 10.5], palette,
             mono_columns=(1, 2, 5), center_columns=(0, 3, 4))
        constraints = []
        for index in table["indexes"]:
            cols = index.get("expressions") or index.get("column_names") or []
            constraints.append((index.get("name"), "唯一索引" if index.get("unique") else "索引",
                                ", ".join(text(c) for c in cols)))
        for unique in table["unique"]:
            constraints.append((unique.get("name"), "唯一约束", ", ".join(unique.get("column_names") or [])))
        for fk in table["foreign_keys"]:
            target = ".".join(v for v in [fk.get("referred_schema"), fk.get("referred_table")] if v)
            detail = f"{', '.join(fk['constrained_columns'])} → {target} ({', '.join(fk['referred_columns'])})"
            if fk.get("options"):
                detail += f"\n{fk['options']}"
            constraints.append((fk.get("name"), "外键", detail))
        for check in table["checks"]:
            constraints.append((check.get("name"), "检查约束", check.get("sqltext")))
        if constraints:
            doc.add_heading("索引与约束", 2)
            grid(doc, ["名称", "种类", "定义"], constraints, [6, 3, 17.7], palette, mono_columns=(0, 2))
    return doc


def save_document(doc, path, force=False):
    """Finish the ZIP before publishing it; refuse accidental replacement."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        doc.save(temporary)
        if force:
            os.replace(temporary, path)
        else:
            os.link(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
